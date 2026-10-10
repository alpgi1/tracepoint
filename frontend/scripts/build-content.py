#!/usr/bin/env python3
"""Build the landing page content file from the repository's own data.

Reads the backend's baseline situation-analysis reports (reports/baseline/*.json), the sales
file and the customer-study file under docs/, and writes frontend/content/markets.json. Every
figure on the map and in the vehicle views comes from this file, with its source named, so the
UI never carries a typed-in number (frontend forbidden list, section 1).

    python3 frontend/scripts/build-content.py

Re-run whenever the baseline reports change. The output is committed so the frontend builds
without the Python toolchain.
"""
from __future__ import annotations

import json
from collections import defaultdict
from datetime import date
from pathlib import Path

import openpyxl
import yaml

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "frontend" / "content" / "markets.json"

MARKETS = [
    {"code": "CN", "name": "China"},
    {"code": "US", "name": "United States"},
    {"code": "EU", "name": "European Union"},
]
SALES_YEAR = "2025"
STUDY_SHEET_CN_EU = "CN_EU_2025"
MIN_ROWS_FOR_DOMAIN = 50  # a domain needs this many classified rows before it can be called the weakest


def load_yaml(path: Path):
    return yaml.safe_load(path.read_text())


def main() -> None:
    vehicles = {v["code"]: v for v in load_yaml(ROOT / "config" / "vehicles.yaml")["vehicles"]}
    domains = load_yaml(ROOT / "config" / "taxonomy" / "domains.yaml")["domains"]

    # Sales per vehicle sheet and market.
    sales_wb = openpyxl.load_workbook(ROOT / "docs" / "sales_volumes.xlsx", data_only=True)
    sales: dict[str, dict[str, dict[str, int]]] = {}
    for ws in sales_wb.worksheets:
        header = [c.value for c in ws[1]]
        for row in ws.iter_rows(min_row=2, values_only=True):
            rec = dict(zip(header, row))
            sales.setdefault(ws.title, {})[rec["market_code"]] = {
                "2024": int(rec["volume_2024"]), "2025": int(rec["volume_2025"]), "2030": int(rec["volume_2030"]),
            }

    # Study means for CN and EU (one column per vehicle).
    study_wb = openpyxl.load_workbook(ROOT / "docs" / "F70_G60_G68_G70_customer_studies.xlsx", data_only=True)
    rows = list(study_wb[STUDY_SHEET_CN_EU].iter_rows(values_only=True))
    columns = {"CN_G68": 1, "CN_G70": 2, "EU_F70": 3, "EU_G60": 4, "EU_G70": 5}
    means: dict[str, dict[str, float]] = {}
    current = None
    for r in rows[2:]:
        if r[0] == "Overall product satisfaction":
            current = r[0]
        elif r[0] and r[0] != "Mean" and r[1] is None:
            current = str(r[0]).strip()
        elif r[0] == "Mean" and current:
            means[current] = {k: float(r[i]) for k, i in columns.items()}

    # Baseline reports: one per vehicle and market with verbatims.
    baselines: dict[tuple[str, str], dict] = {}
    for path in sorted((ROOT / "reports" / "baseline").glob("*.json")):
        d = json.loads(path.read_text())
        baselines[(d["vehicle"]["code"], d["market"])] = d

    def domain_name(key: str) -> str:
        return domains[key]["name"] if key in domains else key

    def weakest_domain(domain_rows: list[dict]) -> dict | None:
        ranked = sorted(
            (x for x in domain_rows if x["domain"] != "general" and x["rows"] >= MIN_ROWS_FOR_DOMAIN),
            key=lambda x: -x["complaint_share"],
        )
        if not ranked:
            return None
        x = ranked[0]
        return {"key": x["domain"], "name": domain_name(x["domain"]), "complaintShare": round(x["complaint_share"], 3), "rows": x["rows"]}

    def weakest_study_attribute(market: str, codes: list[str]) -> dict | None:
        cols = [f"{market}_{c}" for c in codes if f"{market}_{c}" in columns]
        if not cols:
            return None
        avg = {a: sum(v[c] for c in cols) / len(cols) for a, v in means.items() if a != "Overall product satisfaction"}
        attr, score = min(avg.items(), key=lambda t: t[1])
        overall = sum(means["Overall product satisfaction"][c] for c in cols) / len(cols)
        return {"attribute": attr, "score": round(score, 2), "overall": round(overall, 2), "scale": "1-10 mean", "vehicles": codes}

    out_markets = []
    for m in MARKETS:
        code = m["code"]
        sold = []
        for sheet, per_market in sales.items():
            units = per_market.get(code, {}).get(SALES_YEAR, 0)
            if units > 0:
                sold.append((sheet, units))
        sales_total = sum(u for _, u in sold)
        verbatims = sum(d["coverage"]["verbatims"] for (v, mk), d in baselines.items() if mk == code)
        # Weakest domain across the market's baselines, by complaint share.
        agg: dict[str, list[int]] = defaultdict(lambda: [0, 0])
        for (v, mk), d in baselines.items():
            if mk != code:
                continue
            for x in d["domains"]:
                agg[x["domain"]][0] += x["rows"]
                agg[x["domain"]][1] += x["complaints"]
        domain_rows = [{"domain": k, "rows": r, "complaints": c, "complaint_share": (c / r if r else 0)} for k, (r, c) in agg.items()]
        weakest = weakest_domain(domain_rows)
        study_codes = {"CN": ["G68", "G70"], "EU": ["F70", "G60", "G70"]}.get(code, [])
        study = weakest_study_attribute(code, study_codes)

        vehicle_entries = []
        for sheet, units in sold:
            # Sheet G60_G68 carries both derivatives; China sells the long-wheelbase G68.
            vcode = "G68" if (sheet == "G60_G68" and code == "CN") else sheet.split("_")[0]
            base = baselines.get((vcode, code)) or (baselines.get(("G60", code)) if vcode == "G68" else None)
            entry = {
                "code": vcode,
                "name": vehicles[vcode]["name"],
                "series": vehicles[vcode]["series"],
                "drawing": "G60" if vcode == "G68" else vcode,
                "sales": {"year": int(SALES_YEAR), "units": units, "source": f"docs/sales_volumes.xlsx sheet {sheet}"},
                "verbatims": base["coverage"]["verbatims"] if base and base["vehicle"]["code"] == vcode else 0,
                "weakestDomain": weakest_domain(base["domains"]) if base and base["vehicle"]["code"] == vcode else None,
                "domains": [
                    {
                        "key": x["domain"],
                        "name": domain_name(x["domain"]),
                        "rows": x["rows"],
                        "complaints": x["complaints"],
                        "likes": x["likes"],
                        "complaintShare": round(x["complaint_share"], 3),
                    }
                    for x in base["domains"]
                ] if base and base["vehicle"]["code"] == vcode else [],
                "complaintThemes": [
                    {
                        "theme": t["theme"],
                        "domain": t["domain"],
                        "rows": t["rows"],
                        "complaints": t["complaints"],
                        "defects": t["defects"],
                        "difficultToUse": t["difficult_to_use"],
                        "wants": t["wants"],
                    }
                    for t in base["top_complaint_themes"]
                ] if base and base["vehicle"]["code"] == vcode else [],
                "study": None,
            }
            col = f"{code}_{vcode}"
            if col in columns:
                ranked = sorted(((a, v[col]) for a, v in means.items() if a != "Overall product satisfaction"), key=lambda t: t[1])
                entry["study"] = {
                    "sheet": STUDY_SHEET_CN_EU,
                    "overall": round(means["Overall product satisfaction"][col], 2),
                    "weakest": [{"attribute": a, "score": round(s, 2)} for a, s in ranked[:3]],
                    "scale": "1-10 mean",
                }
            elif base and base["studies"]:
                st = base["studies"][0]
                entry["study"] = {
                    "sheet": st["study"],
                    "overall": None,
                    "weakest": [{"attribute": w["attribute"], "score": round(w["score"], 2)} for w in st["weakest"][:3]],
                    "scale": st["weakest"][0]["scale"] if st["weakest"] else None,
                }
            vehicle_entries.append(entry)

        out_markets.append({
            "code": code,
            "name": m["name"],
            "sales": {"year": int(SALES_YEAR), "units": sales_total, "source": "docs/sales_volumes.xlsx, all vehicle sheets"},
            "verbatims": {"count": verbatims, "source": "reports/baseline/*.json coverage.verbatims (unique verbatims per vehicle and market)"},
            "weakestDomain": weakest,
            "weakestDomainSource": "Highest complaint share among domains with at least 50 classified rows, summed over the market's baseline reports",
            "study": study,
            "vehicles": vehicle_entries,
        })

    payload = {
        "generatedAt": date.today().isoformat(),
        "dataVintage": {
            "sales": "sales_volumes.xlsx: 2025 actual, 2030 forecast",
            "verbatims": "feedback_hackathon.xlsx per vehicle, classified rows as ingested by Phase 1",
            "studies": "CN_EU_2025 means (1-10), US_2025 distributions (1-7 derived mean)",
            "baselineReports": sorted(f"{v}_{m}" for v, m in baselines),
        },
        "domains": [{"key": k, "name": v["name"], "description": v["description"]} for k, v in domains.items()],
        "markets": out_markets,
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(payload, indent=2, ensure_ascii=False) + "\n")
    print(f"wrote {OUT.relative_to(ROOT)}: {len(out_markets)} markets, {sum(len(m['vehicles']) for m in out_markets)} vehicle entries")
    for m in out_markets:
        wd = m["weakestDomain"]["name"] if m["weakestDomain"] else (m["study"]["attribute"] if m["study"] else None)
        print(f"  {m['code']}: sales {m['sales']['units']:,} verbatims {m['verbatims']['count']:,} weakest {wd} vehicles {[v['code'] for v in m['vehicles']]}")


if __name__ == "__main__":
    main()
