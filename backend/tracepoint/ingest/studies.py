"""Ingest the customer study workbook (CN/EU attribute means and US rating distributions)."""

from __future__ import annotations

import re
from collections import Counter
from typing import Any

import openpyxl
import yaml
from sqlalchemy import delete
from sqlalchemy.orm import Session

from tracepoint.core.audit import record_event
from tracepoint.core.ids import slug
from tracepoint.core.settings import get_settings
from tracepoint.db.models import Run, StudyMetric
from tracepoint.ingest.registry import file_path, register_source

RATING_LABELS = [
    "I Hate It",
    "A Failure",
    "Unsatisfactory",
    "Satisfactory",
    "Excellent",
    "Delightful",
    "I Love It",
]
RATING_SCORE = {label: i + 1 for i, label in enumerate(RATING_LABELS)}
MODEL_RE = re.compile(r"\b(F70|G60|G68|G70)\b")
MARKET_NAMES = {"china": "CN", "eu": "EU", "europe": "EU", "us": "US", "usa": "US"}


def _attribute_map() -> dict[str, dict[str, Any]]:
    path = get_settings().config_dir / "studies" / "attribute_map.yaml"
    if not path.exists():
        return {}
    with path.open(encoding="utf-8") as fh:
        data = yaml.safe_load(fh) or {}
    return {e["raw"].strip(): e for e in data.get("entries", [])}


def _cell(v: Any) -> str:
    return "" if v is None else str(v).strip()


def _parse_cn_eu(ws, source_id: str, run_id: str, amap: dict, dq: Counter) -> list[StudyMetric]:
    rows = list(ws.iter_rows(values_only=True))
    # row 2 holds the market labels at the first column of each market block
    market_by_col: dict[int, str] = {}
    current = None
    for col, v in enumerate(rows[1]):
        label = _cell(v).lower()
        if label in MARKET_NAMES:
            current = MARKET_NAMES[label]
        if col >= 1 and current:
            market_by_col[col] = current
    # row 3 holds the model names per column (first attribute row)
    model_by_col: dict[int, str] = {}
    for col, v in enumerate(rows[2]):
        m = MODEL_RE.search(_cell(v))
        if m and col >= 1:
            model_by_col[col] = m.group(1)
    metrics: list[StudyMetric] = []
    attribute = None
    attr_row = None
    for r_idx, row in enumerate(rows, start=1):
        first = _cell(row[0]) if row else ""
        if first and first.lower() != "mean":
            attribute = first
            attr_row = r_idx
            continue
        if first.lower() == "mean" and attribute:
            for col, model in model_by_col.items():
                val = row[col] if col < len(row) else None
                if val is None or _cell(val) == "":
                    dq["cn_eu_missing_value"] += 1
                    continue
                market = market_by_col.get(col)
                if market is None:
                    dq["cn_eu_missing_market"] += 1
                    continue
                entry = amap.get(attribute, {})
                key = entry.get("key") or slug(attribute)
                metrics.append(
                    StudyMetric(
                        id=f"MET-CNEU-{market}-{model}-{key}"[:48],
                        vehicle_code=model,
                        market_code=market,
                        study="CN_EU_2025",
                        attribute_raw=attribute,
                        attribute_key=key,
                        domain_key=entry.get("domain"),
                        stat_kind="mean",
                        mean=float(val),
                        distribution=None,
                        scale_note="mean rating; scale not stated in the file (assumed 1-10)",
                        source_id=source_id,
                        locator={
                            "sheet": ws.title,
                            "attribute_row": attr_row,
                            "row": r_idx,
                            "col": col + 1,
                        },
                        created_run_id=run_id,
                    )
                )
            attribute = None
    return metrics


def _parse_us(ws, source_id: str, run_id: str, amap: dict, dq: Counter) -> list[StudyMetric]:
    rows = list(ws.iter_rows(values_only=True))
    model_by_col: dict[int, str] = {}
    for col, v in enumerate(rows[0]):
        m = MODEL_RE.search(_cell(v))
        if m and col >= 2:
            model_by_col[col] = m.group(1)
    metrics: list[StudyMetric] = []
    attribute = None
    attr_row = None
    dist: dict[str, dict[str, float]] = {}
    dist_rows: dict[str, list[int]] = {}

    def flush() -> None:
        nonlocal attribute, dist, dist_rows
        if attribute is None:
            return
        for col, model in model_by_col.items():
            d = dist.get(model, {})
            if len(d) != len(RATING_LABELS):
                dq["us_incomplete_distribution"] += 1
                continue
            total = sum(d.values())
            if abs(total - 1.0) > 0.02:
                dq["us_distribution_sum_off"] += 1
            mean_1_7 = sum(RATING_SCORE[k] * v for k, v in d.items()) / total if total else None
            payload = {
                "shares": d,
                "scale": "7-step verbal scale, share of respondents (sample base = 1)",
                "mean_1_7": round(mean_1_7, 4) if mean_1_7 is not None else None,
                "top2": round(d["Delightful"] + d["I Love It"], 4),
                "bottom3": round(d["I Hate It"] + d["A Failure"] + d["Unsatisfactory"], 4),
            }
            entry = amap.get(attribute, {})
            key = entry.get("key") or slug(attribute)
            metrics.append(
                StudyMetric(
                    id=f"MET-US-US-{model}-{key}"[:48],
                    vehicle_code=model,
                    market_code="US",
                    study="US_2025",
                    attribute_raw=attribute,
                    attribute_key=key,
                    domain_key=entry.get("domain"),
                    stat_kind="distribution",
                    mean=payload["mean_1_7"],
                    distribution=payload,
                    scale_note="share of respondents per rating label; derived mean on 1-7 scale",
                    source_id=source_id,
                    locator={
                        "sheet": ws.title,
                        "attribute_row": attr_row,
                        "rows": dist_rows.get(model, []),
                        "col": col + 1,
                    },
                    created_run_id=run_id,
                )
            )
        attribute = None
        dist = {}
        dist_rows = {}

    for r_idx, row in enumerate(rows, start=1):
        first = _cell(row[0]) if row else ""
        second = _cell(row[1]) if row and len(row) > 1 else ""
        if not first:
            continue
        if first in RATING_SCORE:
            for col, model in model_by_col.items():
                val = row[col] if col < len(row) else None
                if val is None or _cell(val) == "":
                    continue
                dist.setdefault(model, {})[first] = float(val)
                dist_rows.setdefault(model, []).append(r_idx)
            continue
        if first.lower().startswith("sample total"):
            continue
        if second == "":
            flush()
            attribute = first
            attr_row = r_idx
    flush()
    return metrics


def ingest_studies(session: Session, run: Run) -> dict[str, int]:
    source = register_source(session, run, "studies")
    path = file_path("studies")
    amap = _attribute_map()
    dq: Counter = Counter()
    session.execute(delete(StudyMetric))
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    metrics = _parse_cn_eu(wb["CN_EU_2025"], source.id, run.id, amap, dq)
    metrics += _parse_us(wb["US_2025"], source.id, run.id, amap, dq)
    unmapped = sum(1 for m in metrics if m.attribute_raw not in amap)
    session.add_all(metrics)
    session.flush()
    stats = {
        "metrics": len(metrics),
        "cn_eu_means": sum(1 for m in metrics if m.study == "CN_EU_2025"),
        "us_distributions": sum(1 for m in metrics if m.study == "US_2025"),
        "unmapped_attributes": unmapped,
        **{f"dq_{k}": v for k, v in dq.items()},
    }
    record_event(
        session,
        actor_type="system",
        actor_id="ingest.studies",
        action="studies.ingested",
        entity_type="source",
        entity_id=source.id,
        run_id=run.id,
        after=stats,
        rationale=f"Parsed {path.name}: CN_EU_2025 means and US_2025 distributions",
    )
    return stats
